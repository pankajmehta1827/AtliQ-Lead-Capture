"""CRM: the system of record for every lead (open, won, lost).

The AI Lead app writes here only through confirmed drafts (services.review.confirm_draft); this router
adds the CRM-side views: full lead list with where each record came from, manual add, CSV/Excel import
and export, and an integration summary of what the AI wrote and when."""
from __future__ import annotations

import csv
import io
from datetime import date, datetime, timezone

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import func, select
from sqlalchemy.orm import Session, joinedload

from .. import serializers as ser
from ..audit import log_action
from ..catalogue import STAGES
from ..db import get_db
from ..deps import current_user
from ..matching import find_matches
from ..models import Deal, DealActivity, Draft
from ..services.seed import CRM_COLUMNS, upsert_crm_rows

router = APIRouter(prefix="/api/crm", tags=["crm"])


def _origins(db: Session) -> dict[int, dict]:
    """Per deal: how it entered the CRM and how often the AI Lead app updated it."""
    out: dict[int, dict] = {}
    for a in db.scalars(select(DealActivity).order_by(DealActivity.created_at)).all():
        info = out.setdefault(a.deal_id, {"origin": "crm_export", "ai_updates": 0, "last_ai_update": None})
        created = isinstance(a.changes, dict) and "created" in a.changes
        if a.draft_id:
            if created:
                info["origin"] = "ai_capture"
            else:
                info["ai_updates"] += 1
            info["last_ai_update"] = a.created_at.isoformat() if a.created_at else None
        elif created:
            info["origin"] = "import" if (a.summary or "").startswith("Created by import") else "manual"
    return out


@router.get("/leads")
def list_leads(db: Session = Depends(get_db), _: str = Depends(current_user)):
    deals = db.scalars(
        select(Deal).options(joinedload(Deal.company), joinedload(Deal.contact)).order_by(Deal.lead_code)
    ).unique().all()
    origins = _origins(db)
    pending = dict(
        db.execute(select(Draft.deal_id, func.count()).where(Draft.status == "pending", Draft.deal_id.is_not(None))
                   .group_by(Draft.deal_id)).all()
    )
    default = {"origin": "crm_export", "ai_updates": 0, "last_ai_update": None}
    return [{**ser.deal(d), **origins.get(d.id, default), "pending_ai_drafts": pending.get(d.id, 0)} for d in deals]


class NewLead(BaseModel):
    company: str = Field(min_length=2, max_length=200)
    contact_name: str | None = None
    contact_email: str | None = None
    source: str | None = None
    service_interest: str | None = None
    status: str | None = "New"
    est_value_usd: float | None = None
    owner: str | None = None
    next_step: str | None = None
    next_followup_date: date | None = None
    notes: str | None = None
    force: bool = False  # create even if it looks like a duplicate


@router.post("/leads")
def create_lead(body: NewLead, user: str = Depends(current_user), db: Session = Depends(get_db)):
    if body.status and body.status not in STAGES:
        raise HTTPException(422, f"status must be one of {STAGES}")
    if not body.force:
        dupes = [m for m in find_matches(db, company=body.company, emails=[body.contact_email or ""]) if m["score"] >= 0.9]
        if dupes:
            raise HTTPException(409, {"message": "This looks like an existing lead", "matches": dupes[:3]})
    result = upsert_crm_rows(
        db,
        [{
            "company": body.company, "contact_name": body.contact_name, "contact_email": body.contact_email,
            "source": body.source, "service_interest": body.service_interest, "status": body.status,
            "est_value_usd": body.est_value_usd, "owner": body.owner or user, "created_date": date.today().isoformat(),
            "next_followup_date": body.next_followup_date, "notes": body.notes,
        }],
        actor=user, source_name="manual", update_existing=False,
    )
    deal = db.scalars(select(Deal).order_by(Deal.id.desc()).limit(1)).first()
    if not result["created"] or not deal:
        raise HTTPException(400, "Lead not created")
    deal.next_step = body.next_step
    act = db.scalar(select(DealActivity).where(DealActivity.deal_id == deal.id))
    if act:
        act.summary = "Created manually"
    log_action(db, actor=user, action="lead_created", outcome="manual", entity_type="deal", entity_id=deal.id,
               details={"lead_code": deal.lead_code})
    db.commit()
    db.refresh(deal)
    return ser.deal(deal)


def _export_rows(db: Session) -> list[list]:
    deals = db.scalars(
        select(Deal).options(joinedload(Deal.company), joinedload(Deal.contact)).order_by(Deal.lead_code)
    ).unique().all()
    rows = []
    for d in deals:
        rows.append([
            d.lead_code, d.company.name if d.company else "", d.contact.name if d.contact and d.contact.name else "",
            d.contact.email if d.contact and d.contact.email else "", d.source or "", d.service_interest or "",
            d.status or "", int(d.est_value_usd) if d.est_value_usd else "", d.owner or "",
            d.created_date.isoformat() if d.created_date else "",
            d.last_contact_date.isoformat() if d.last_contact_date else "",
            d.next_followup_date.isoformat() if d.next_followup_date else "",
            d.notes or "",
        ])
    return rows


@router.get("/export")
def export_leads(format: str = "csv", user: str = Depends(current_user), db: Session = Depends(get_db)):
    rows = _export_rows(db)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d")
    log_action(db, actor=user, action="crm_export", outcome=format, details={"rows": len(rows)})
    db.commit()
    if format == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font

        wb = Workbook()
        ws = wb.active
        ws.title = "CRM"
        ws.append(CRM_COLUMNS)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append(r)
        ws.freeze_panes = "A2"
        for col, width in zip("ABCDEFGHIJKLM", [9, 26, 22, 30, 14, 20, 14, 13, 11, 12, 15, 17, 60]):
            ws.column_dimensions[col].width = width
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return StreamingResponse(
            buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="atliq_crm_{stamp}.xlsx"'},
        )
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(CRM_COLUMNS)
    w.writerows(rows)
    return StreamingResponse(
        iter([out.getvalue()]), media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="atliq_crm_{stamp}.csv"'},
    )


@router.post("/import")
async def import_leads(
    file: UploadFile = File(...),
    update_existing: bool = Form(True),
    user: str = Depends(current_user),
    db: Session = Depends(get_db),
):
    name = file.filename or "upload"
    data = await file.read()
    if len(data) > 5_000_000:
        raise HTTPException(413, "File too large (max 5 MB)")
    if name.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        try:
            ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        except Exception as exc:  # noqa: BLE001
            raise HTTPException(422, f"Could not read Excel file: {exc}") from exc
        it = ws.iter_rows(values_only=True)
        header = [str(h).strip().lower() if h is not None else "" for h in next(it, [])]
        rows = [dict(zip(header, r)) for r in it if any(v not in (None, "") for v in r)]
    elif name.lower().endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        rows = list(csv.DictReader(io.StringIO(text)))
        header = [h.strip().lower() for h in (rows[0].keys() if rows else [])]
    else:
        raise HTTPException(422, "Upload a .csv or .xlsx file")
    if rows and "company" not in header:
        raise HTTPException(422, f"Missing 'company' column. Expected columns: {', '.join(CRM_COLUMNS)}")
    return upsert_crm_rows(db, rows, user, name, update_existing)


@router.get("/sync")
def sync_status(db: Session = Depends(get_db), _: str = Depends(current_user)):
    """Integration view: what the AI Lead app has written into the CRM (always via an owner's confirmation)."""
    acts = db.scalars(
        select(DealActivity).options(joinedload(DealActivity.deal).joinedload(Deal.company), joinedload(DealActivity.source_item))
        .where(DealActivity.draft_id.is_not(None)).order_by(DealActivity.created_at.desc())
    ).unique().all()
    created = {a.deal_id for a in acts if isinstance(a.changes, dict) and "created" in a.changes}
    updated = {a.deal_id for a in acts} - created
    pending = db.scalar(select(func.count()).select_from(Draft).where(Draft.status == "pending")) or 0
    return {
        "ai_created": len(created),
        "ai_updated": len(updated),
        "writes": len(acts),
        "last_sync": acts[0].created_at.isoformat() if acts else None,
        "pending_drafts": pending,
        "recent": [
            {
                "deal_id": a.deal_id, "lead_code": a.deal.lead_code if a.deal else None,
                "company": a.deal.company.name if a.deal and a.deal.company else None,
                "kind": "created" if a.deal_id in created and "created" in (a.changes or {}) else "updated",
                "fields": [k for k in (a.changes or {}) if k != "created"],
                "actor": a.actor, "at": a.created_at.isoformat() if a.created_at else None,
                "source": a.source_item.subject if a.source_item else None,
                "channel": a.source_item.channel if a.source_item else None,
            }
            for a in acts[:12]
        ],
    }
